import json
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "work" / "raw"
ANALYSIS = ROOT / "work" / "analysis"
OUT = ROOT / "outputs" / "geo_ai_jiangmen_20260729"

PLATFORMS = {
    "deepseek": "DeepSeek",
    "doubao": "豆包",
    "yuanbao": "腾讯元宝",
    "wenxin": "文心一言",
    "qianwen": "千问",
}

raw_audit = {}
all_valid = True
for folder, name in PLATFORMS.items():
    files = sorted((RAW / folder).glob("q*.json"))
    numbers = []
    invalid = []
    for file in files:
        try:
            data = json.loads(file.read_text(encoding="utf-8"))
            number = int(file.stem[1:])
            numbers.append(number)
            messages = data.get("messages", [])
            answer = messages[-1] if len(messages) >= 2 else ""
            if not str(answer).strip():
                invalid.append(f"{file.name}: empty answer message")
            if int(data.get("question_number", -1)) != number:
                invalid.append(f"{file.name}: question number mismatch")
        except Exception as exc:
            invalid.append(f"{file.name}: {type(exc).__name__}")
    missing = sorted(set(range(1, 37)) - set(numbers))
    valid = len(files) == 36 and not missing and not invalid
    all_valid = all_valid and valid
    raw_audit[name] = {
        "files": len(files),
        "missing": missing,
        "invalid": invalid,
        "valid": valid,
    }

rows = json.loads((ANALYSIS / "recommendations_enriched.json").read_text(encoding="utf-8"))
summary = json.loads((ANALYSIS / "final_summary.json").read_text(encoding="utf-8"))
support = Counter(
    "有效" if x["support"].startswith("是（")
    else "部分支持" if x["support"].startswith("部分支持")
    else "不支持" if x["support"].startswith("否（")
    else "未核验"
    for x in rows
)

xlsx_path = OUT / "江门AI搜索可见性测试明细_180题.xlsx"
docx_path = OUT / "江门全屋定制、装修与全屋设计AI搜索基准报告.docx"
pdf_path = OUT / "report_render" / "江门全屋定制、装修与全屋设计AI搜索基准报告.pdf"

with zipfile.ZipFile(xlsx_path) as zf:
    xlsx_zip_error = zf.testzip()
with zipfile.ZipFile(docx_path) as zf:
    docx_zip_error = zf.testzip()

workbook = openpyxl.load_workbook(xlsx_path, read_only=False, data_only=False)
report = Document(docx_path)
report_text = "\n".join(p.text for p in report.paragraphs)

required_report_sections = [
    "执行摘要",
    "研究设计与证据口径",
    "总体品牌可见性",
    "五条关键词赛道 Top10",
    "平台差异",
    "引用网页与来源质量",
    "官网依赖与第三方来源",
    "楼盘场景",
    "江门本地品牌 vs 全国连锁品牌",
    "当前GEO竞争最强的品牌",
    "公开内容的明显空白选题",
    "结论",
]

audit = {
    "raw_samples": raw_audit,
    "raw_sample_total": sum(x["files"] for x in raw_audit.values()),
    "raw_samples_valid": all_valid,
    "recommendation_rows": len(rows),
    "unique_brands": len({x["brand"] for x in rows}),
    "support_counts": dict(support),
    "candidate_pages": summary["candidate_pages"],
    "opened_candidate_pages": summary["opened_candidate_pages"],
    "unique_citations": summary["unique_citation_count"],
    "workbook": {
        "exists": xlsx_path.exists(),
        "size": xlsx_path.stat().st_size,
        "zip_error": xlsx_zip_error,
        "sheets": workbook.sheetnames,
        "detail_rows_including_title_and_header": workbook["180题明细"].max_row,
        "brand_rows_including_title_and_header": workbook["品牌统计"].max_row,
        "citation_rows_including_title_and_header": workbook["引用核验"].max_row,
        "formula_count": sum(
            1
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ),
    },
    "report": {
        "exists": docx_path.exists(),
        "size": docx_path.stat().st_size,
        "zip_error": docx_zip_error,
        "paragraphs": len(report.paragraphs),
        "tables": len(report.tables),
        "missing_sections": [x for x in required_report_sections if x not in report_text],
        "pdf_exists": pdf_path.exists(),
        "pdf_size": pdf_path.stat().st_size if pdf_path.exists() else 0,
    },
}

audit["complete"] = all([
    audit["raw_samples_valid"],
    audit["raw_sample_total"] == 180,
    audit["recommendation_rows"] == 757,
    audit["unique_brands"] == 105,
    audit["candidate_pages"] == 135,
    audit["opened_candidate_pages"] >= 120,
    audit["workbook"]["exists"],
    audit["workbook"]["zip_error"] is None,
    audit["workbook"]["detail_rows_including_title_and_header"] == 760,
    len(audit["workbook"]["sheets"]) == 9,
    audit["workbook"]["formula_count"] >= 1,
    audit["report"]["exists"],
    audit["report"]["zip_error"] is None,
    not audit["report"]["missing_sections"],
    audit["report"]["pdf_exists"],
])

(ANALYSIS / "completion_audit.json").write_text(
    json.dumps(audit, ensure_ascii=False, indent=2),
    encoding="utf-8",
)
print(json.dumps(audit, ensure_ascii=False, indent=2))
if not audit["complete"]:
    raise SystemExit(1)
