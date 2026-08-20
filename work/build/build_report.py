import json
import math
import re
from collections import Counter
from pathlib import Path
from urllib.parse import urlparse

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
ANALYSIS = ROOT / "work" / "analysis"
OUT = ROOT / "outputs" / "geo_ai_jiangmen_20260729"
OUT.mkdir(parents=True, exist_ok=True)

summary = json.loads((ANALYSIS / "final_summary.json").read_text(encoding="utf-8"))
rows = json.loads((ANALYSIS / "recommendations_enriched.json").read_text(encoding="utf-8"))

NAVY = "17324D"
BLUE = "2E6F95"
TEAL = "2A9D8F"
SAND = "E9C46A"
CORAL = "E76F51"
INK = "23313F"
PALE = "EEF4F7"
LINE = "D7E1E8"
WHITE = "FFFFFF"
MUTED = "607080"

FONT_CN = Path("C:/Windows/Fonts/msyh.ttc")
FONT_BOLD_CN = Path("C:/Windows/Fonts/msyhbd.ttc")
def image_font(size, bold=False):
    path = FONT_BOLD_CN if bold and FONT_BOLD_CN.exists() else FONT_CN
    if path.exists():
        return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_header(section, text):
    header = section.header
    p = header.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(8.5)
    run.font.color.rgb = rgb(MUTED)
    p.paragraph_format.space_after = Pt(0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(MUTED)


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    r_pr.append(color_node)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:eastAsia"), "微软雅黑")
    r_pr.append(r_fonts)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_external_file_link(paragraph, text, file_path):
    # Use a relative target so the link stays valid when the output folder is moved as a unit.
    return add_hyperlink(paragraph, text, file_path.name)


def set_run_font(run, size=None, bold=None, color=None, font_name="Microsoft YaHei"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = rgb(color)


def add_body(doc, text="", bold_lead=None, color=INK, size=10.5, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(space_after)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, size=size, bold=True, color=color)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=size, color=color)
    else:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color)
    return p


def add_bullet(doc, text, level=0, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.25 + 0.2 * level)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.18
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=10, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10, color=INK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = doc.styles[f"Heading {level}"]
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_font(
        run,
        size=17 if level == 1 else 12.5,
        bold=True,
        color=NAVY if level == 1 else BLUE,
    )
    return p


def add_kicker(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    set_run_font(r, size=8.5, bold=True, color=TEAL)
    return p


def add_callout(doc, title, body, fill=PALE, accent=TEAL):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.12)
    table.columns[1].width = Inches(6.15)
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_margins(table.cell(0, 1), top=130, bottom=130, start=150, end=150)
    p = table.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    p2 = table.cell(0, 1).add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.18
    r2 = p2.add_run(body)
    set_run_font(r2, size=9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, data, widths=None, font_size=8.5, header_fill=BLUE, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(str(value))
        set_run_font(r, size=font_size, bold=True, color=WHITE)
    if widths:
        for i, width in enumerate(widths):
            for cell in table.columns[i].cells:
                cell.width = Inches(width)
    for row_index, values in enumerate(data):
        row = table.add_row()
        prevent_row_split(row)
        if row_index % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, "F7FAFC")
        for i, value in enumerate(values):
            cell = row.cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run("" if value is None else str(value))
            set_run_font(r, size=font_size, bold=first_col_bold and i == 0, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_image(doc, path, width=6.3, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(7)
        r = cp.add_run(caption)
        set_run_font(r, size=8, color=MUTED)


def pct(n, d):
    return f"{n / d * 100:.1f}%" if d else "—"


def make_charts():
    top = summary["overall"][:10]
    p1 = OUT / "chart_overall_top10.png"
    im = Image.new("RGB", (1500, 780), "white")
    d = ImageDraw.Draw(im)
    d.text((60, 32), "总体AI推荐次数 TOP10", font=image_font(38, True), fill=f"#{NAVY}")
    max_count = max(x["count"] for x in top)
    x0, x1 = 310, 1400
    y0, step = 115, 60
    for i, item in enumerate(top):
        y = y0 + i * step
        d.text((60, y + 9), item["brand"], font=image_font(24), fill=f"#{INK}")
        width = int((x1 - x0) * item["count"] / max_count)
        d.rounded_rectangle((x0, y, x0 + width, y + 36), radius=8, fill=f"#{BLUE}")
        d.text((x0 + width + 12, y + 4), str(item["count"]), font=image_font(23, True), fill=f"#{NAVY}")
    im.save(p1)

    labels = ["有效支持", "部分支持", "不支持", "未核验"]
    values = [192, 31, 119, 415]
    colors = ["#2A9D8F", "#E9C46A", "#E76F51", "#A9B8C5"]
    p2 = OUT / "chart_evidence.png"
    im = Image.new("RGB", (1200, 680), "white")
    d = ImageDraw.Draw(im)
    d.text((55, 30), "引用证据分级", font=image_font(38, True), fill=f"#{NAVY}")
    box = (70, 110, 610, 650)
    start = -90
    total = sum(values)
    for value, color in zip(values, colors):
        end = start + 360 * value / total
        d.pieslice(box, start=start, end=end, fill=color, outline="white", width=3)
        start = end
    inner = (210, 250, 470, 510)
    d.ellipse(inner, fill="white")
    d.text((280, 325), "757", font=image_font(52, True), fill=f"#{NAVY}")
    d.text((275, 390), "推荐记录", font=image_font(23), fill=f"#{MUTED}")
    for i, (label, value, color) in enumerate(zip(labels, values, colors)):
        y = 190 + i * 92
        d.rounded_rectangle((720, y, 760, y + 40), radius=7, fill=color)
        d.text((785, y - 1), f"{label}  {value}", font=image_font(28), fill=f"#{INK}")
    im.save(p2)

    source_items = [x for x in summary["source_type_stats"] if x["source_type"] != "未提供链接"]
    p3 = OUT / "chart_sources.png"
    im = Image.new("RGB", (1500, 840), "white")
    d = ImageDraw.Draw(im)
    d.text((55, 30), "不同来源类型的证据结构", font=image_font(38, True), fill=f"#{NAVY}")
    d.rounded_rectangle((860, 45, 900, 82), radius=6, fill=f"#{TEAL}")
    d.text((915, 46), "有效或部分支持", font=image_font(21), fill=f"#{INK}")
    d.rounded_rectangle((1165, 45, 1205, 82), radius=6, fill="#CBD5DC")
    d.text((1220, 46), "未核验或不支持", font=image_font(21), fill=f"#{INK}")
    max_total = max(x["recommendation_rows"] for x in source_items)
    x0, x1 = 330, 1420
    for i, item in enumerate(source_items):
        y = 125 + i * 72
        d.text((55, y + 8), item["source_type"], font=image_font(24), fill=f"#{INK}")
        good = item["valid"] + item["partial"]
        bad = item["unsupported"] + item["unverified"]
        w_good = int((x1 - x0) * good / max_total)
        w_bad = int((x1 - x0) * bad / max_total)
        d.rectangle((x0, y, x0 + w_good, y + 40), fill=f"#{TEAL}")
        d.rectangle((x0 + w_good, y, x0 + w_good + w_bad, y + 40), fill="#CBD5DC")
        d.text((x0 + w_good + w_bad + 10, y + 5), str(good + bad), font=image_font(21, True), fill=f"#{NAVY}")
    im.save(p3)

    scopes = {x["scope"]: x for x in summary["scope_stats"]}
    scope_names = ["全国/跨区域品牌", "江门本地品牌", "其他/未判定"]
    vals = [scopes[x]["recommendation_rows"] for x in scope_names]
    colors = ["#2E6F95", "#2A9D8F", "#A9B8C5"]
    p4 = OUT / "chart_scope.png"
    im = Image.new("RGB", (1320, 650), "white")
    d = ImageDraw.Draw(im)
    d.text((55, 30), "本地与全国品牌的AI推荐量", font=image_font(38, True), fill=f"#{NAVY}")
    base_y = 555
    max_val = max(vals)
    for i, (name, value, color) in enumerate(zip(scope_names, vals, colors)):
        x = 130 + i * 400
        h = int(390 * value / max_val)
        d.rounded_rectangle((x, base_y - h, x + 230, base_y), radius=10, fill=color)
        d.text((x + 58, base_y - h - 72), str(value), font=image_font(36, True), fill=f"#{NAVY}")
        d.text((x + 57, base_y - h - 34), pct(value, 757), font=image_font(22), fill=f"#{MUTED}")
        d.text((x, base_y + 20), name, font=image_font(24), fill=f"#{INK}")
    im.save(p4)
    return p1, p2, p3, p4


chart_overall, chart_evidence, chart_sources, chart_scope = make_charts()

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = rgb(INK)
styles["Normal"].paragraph_format.line_spacing = 1.22
styles["Normal"].paragraph_format.space_after = Pt(5)
for name in ("Heading 1", "Heading 2", "Heading 3"):
    styles[name].font.name = "Microsoft YaHei"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

# Editorial cover: a compact top band, strong title, research metadata, and a single evidence-led statement.
cover_band = doc.add_table(rows=1, cols=1)
cover_band.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = cover_band.cell(0, 0)
set_cell_shading(cell, NAVY)
set_cell_margins(cell, top=260, bottom=260, start=260, end=260)
p = cell.paragraphs[0]
r = p.add_run("GEO / GENERATIVE ENGINE OPTIMIZATION")
set_run_font(r, size=9, bold=True, color=SAND)
p2 = cell.add_paragraph()
p2.paragraph_format.space_before = Pt(6)
r2 = p2.add_run("江门全屋定制、装修与全屋设计\nAI搜索基准报告")
set_run_font(r2, size=27, bold=True, color=WHITE)
p3 = cell.add_paragraph()
p3.paragraph_format.space_before = Pt(9)
r3 = p3.add_run("五平台 × 36问题 × 180次独立联网盲测")
set_run_font(r3, size=12, bold=True, color="DDE8EE")

doc.add_paragraph().paragraph_format.space_after = Pt(16)
add_kicker(doc, "Research baseline / 2026")
cover_intro = doc.add_paragraph()
cover_intro.paragraph_format.space_after = Pt(12)
cover_intro.paragraph_format.line_spacing = 1.12
rr = cover_intro.add_run("公开网络如何塑造AI对江门家居品牌的推荐")
set_run_font(rr, size=18, bold=True, color=NAVY)
add_body(
    doc,
    "本报告基于DeepSeek、豆包、腾讯元宝、文心一言和千问的实时联网结果，逐题记录推荐品牌、推荐位次、理由、引用页面与证据有效性。测试为无品牌盲测，不把任何预设品牌写入提示词。",
    size=11,
    space_after=13,
)
meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
meta.columns[0].width = Inches(1.35)
meta.columns[1].width = Inches(5.0)
for i, (label, value) in enumerate([
    ("搜索期间", "2026-07-25—2026-07-29"),
    ("报告日期", "2026-07-29"),
    ("样本规模", "180次独立测试；757条品牌推荐记录"),
    ("证据核验", "202个去重引用URL；135个候选页实际访问核验"),
]):
    set_cell_shading(meta.cell(i, 0), PALE)
    set_cell_margins(meta.cell(i, 0))
    set_cell_margins(meta.cell(i, 1))
    p = meta.cell(i, 0).paragraphs[0]
    r = p.add_run(label)
    set_run_font(r, size=9, bold=True, color=BLUE)
    p = meta.cell(i, 1).paragraphs[0]
    r = p.add_run(value)
    set_run_font(r, size=9.5, color=INK)
add_callout(
    doc,
    "重要声明",
    "文中“排名”“第一名”“Top10”均仅表示本次AI搜索输出中的可见性，不是官方市场排名、销量排名或质量排名。",
    fill="FFF8E7",
    accent=SAND,
)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

section.different_first_page_header_footer = True
set_repeat_header(section, "江门家居、装修与全屋定制 AI 搜索基准报告｜2026-07-29")
add_page_number(section.footer.paragraphs[0])

add_kicker(doc, "Executive summary")
add_heading(doc, "执行摘要", 1)
add_callout(
    doc,
    "核心结论",
    "江门AI推荐可见性已形成两类优势：全国/跨区域品牌凭借稳定的官网、区域站和广泛第三方内容获得更高覆盖；江门本地品牌在高端定制和细分整装中能取得更靠前位次，但证据来源更集中于公众号、软文或单篇媒体稿，结构化官网案例不足。",
)
add_body(
    doc,
    "总体推荐次数最高的品牌依次为华浔品味装饰（72次）、星艺装饰（64次）、名匠装饰（56次）、健威家居（51次）、劳卡（50次）、索菲亚（48次）和欧派（47次）。其中健威家居平均位次1.69，是总体Top10中平均位置最靠前者；华浔、星艺、名匠则在装修、高端装修和全屋设计三条赛道形成跨题覆盖。",
)
add_image(doc, chart_overall, width=6.2, caption="图1　本次180次AI搜索的总体推荐次数Top10")
add_bullet(doc, "全屋定制：索菲亚22次、欧派21次、健威家居15次居前。", bold_lead="全屋定制：")
add_bullet(doc, "装修：星艺装饰21次，名匠装饰与华浔品味装饰均20次。", bold_lead="装修：")
add_bullet(doc, "高端定制：健威家居13次、盛世周木匠12次，江门本地品牌优势最明显。", bold_lead="高端定制：")
add_bullet(doc, "高端装修：名匠装饰22次、华浔品味装饰21次、星艺装饰19次。", bold_lead="高端装修：")
add_bullet(doc, "全屋设计：华浔品味装饰17次、星艺装饰16次、劳卡12次。", bold_lead="全屋设计：")
add_bullet(doc, "楼盘场景：欧派11次，索菲亚和华浔品味装饰各10次。", bold_lead="楼盘场景：")
add_body(
    doc,
    "证据质量并不与推荐频次等同。757条推荐中，192条引用页面可同时核验品牌、江门与相关业务；31条只能部分支持；119条页面成功打开但未检出对应品牌；415条因无链接、引用与品牌未建立对应或页面不可达保持未核验。",
)
add_image(doc, chart_evidence, width=5.6, caption="图2　推荐记录的引用证据分级")

add_heading(doc, "研究设计与证据口径", 1)
add_heading(doc, "1.1 测试设计", 2)
method_rows = [
    ["平台", "DeepSeek、豆包、腾讯元宝、文心一言、千问"],
    ["问题", "36个问题，覆盖5条关键词赛道与6个楼盘场景"],
    ["独立性", "每题独立开启搜索/会话，不把上一题答案带入下一题"],
    ["推荐记录", "记录AI明确推荐的3—5个品牌；组合推荐拆分记录"],
    ["时间", "2026-07-25—2026-07-29，Asia/Shanghai"],
    ["盲测原则", "提示词不预置任何品牌；保留AI输出的异常和相关性错误"],
]
add_table(doc, ["要素", "执行方式"], method_rows, widths=[1.25, 5.15], font_size=9)
add_heading(doc, "1.2 引用核验规则", 2)
add_bullet(doc, "搜索结果摘要不直接作为证据；先取出AI引用的原始URL，再打开页面核验。")
add_bullet(doc, "“有效支持”要求页面成功打开并同时检出推荐品牌、“江门”和相关业务词。")
add_bullet(doc, "“部分支持”表示只验证了品牌身份、品牌+业务或品牌+江门，不能完整支撑推荐理由。")
add_bullet(doc, "页面打开但没有对应品牌，标为“不支持”；不能访问或没有链接，标为“未核验”。")
add_body(
    doc,
    "本方法是保守口径：它能判断页面是否与推荐结论有基本对应关系，但不把品牌官网自述自动等同于真实交付质量，也不对无法访问的页面作推测。",
)

add_heading(doc, "总体品牌可见性", 1)
top20 = summary["overall"][:20]
top20_rows = []
for i, x in enumerate(top20, 1):
    top20_rows.append([
        i,
        x["brand"],
        x["count"],
        f'{x["average_rank"]:.2f}',
        x["main_tracks"][0]["track"] if x["main_tracks"] else "",
        x["brand_scope"],
    ])
add_table(
    doc,
    ["位次", "品牌", "推荐次数", "平均位次", "主要关联关键词", "范围"],
    top20_rows,
    widths=[0.45, 1.35, 0.75, 0.75, 1.65, 1.2],
    font_size=7.8,
)
add_body(
    doc,
    "完整105个品牌的推荐次数、平均位次、主要关键词、最常引用网站、官网占比及证据支持数，见配套工作簿“品牌统计”表。",
    size=9,
    color=MUTED,
)
add_heading(doc, "2.1 可见性结构", 2)
add_bullet(doc, "装修/设计集团形成跨赛道优势：华浔、星艺、名匠在“装修—高端装修—全屋设计”连续出现，带来高总频次。")
add_bullet(doc, "定制头部全国品牌集中占据基础需求：索菲亚、欧派在全屋定制和楼盘场景中覆盖最广。")
add_bullet(doc, "江门本地品牌更容易在高端、整木、红木语义中被AI选中：健威、盛世周木匠、斯米帝、森润整木表现突出。")
add_bullet(doc, "部分品牌的高频依赖单一页面：峰尚汇28次推荐中，20次引用同一篇中华网页面，内容集中度明显。")

add_heading(doc, "五条关键词赛道 Top10", 1)
track_descriptions = {
    "江门全屋定制": "全国定制品牌占据前列，本地品牌健威、斯米帝、杰普、盛世周木匠进入Top10。",
    "江门装修": "星艺、名匠、华浔形成第一梯队；本地峰尚汇、森之原进入中上位。",
    "江门高端定制品牌": "本地品牌优势最强，健威与盛世周木匠分列前两位。",
    "江门高端装修": "名匠、华浔、星艺高度集中；本地森之原平均位次1.86。",
    "江门全屋设计": "华浔、星艺、劳卡覆盖“设计—施工—定制—交付”的一体化表达。",
}
for track in list(track_descriptions):
    add_heading(doc, track, 2)
    add_body(doc, track_descriptions[track], size=9.5)
    data = [
        [i + 1, x["brand"], x["count"], f'{x["average_rank"]:.2f}', x["brand_scope"]]
        for i, x in enumerate(summary["track_top10"][track])
    ]
    add_table(
        doc,
        ["赛道位次", "品牌", "推荐次数", "平均位次", "品牌范围"],
        data,
        widths=[0.75, 1.65, 0.9, 0.9, 1.55],
        font_size=8.5,
    )

add_heading(doc, "平台差异", 1)
platform_rows = []
for x in summary["platform_stats"]:
    platform_rows.append([
        x["platform"],
        x["sample_count"],
        x["recommendation_rows"],
        x["unique_brands"],
        x["linked_rows"],
        x["valid"] + x["partial"],
        x["unsupported"],
        x["unverified"],
    ])
add_table(
    doc,
    ["平台", "测试", "推荐记录", "品牌数", "带链接", "有效/部分", "不支持", "未核验"],
    platform_rows,
    widths=[1.05, 0.55, 0.8, 0.6, 0.7, 0.8, 0.7, 0.7],
    font_size=8,
)
add_bullet(doc, "DeepSeek给出141条推荐，但保存样本中没有可打开引用链接，因此全部保持未核验。")
add_bullet(doc, "豆包给出168条推荐且全部带链接，99条达到有效或部分支持，是本次证据可复核性最高的平台。")
add_bullet(doc, "千问第29、31题明确没有给出品牌推荐；其第18题出现与高端定制语义明显不符的门窗、摩托、石材等品牌，已原样保留为相关性失误样本。")
add_bullet(doc, "元宝和文心引用量较高，但存在将同一页面分配给多个无关品牌的现象，必须逐品牌核验，不能只看引用数量。")

add_heading(doc, "引用网页与来源质量", 1)
add_image(doc, chart_sources, width=6.2, caption="图3　来源类型的有效/部分支持与未核验/不支持结构")
source_rows = [
    [
        x["source_type"],
        x["recommendation_rows"],
        x["unique_urls"],
        x["valid"],
        x["partial"],
        x["unsupported"],
        x["unverified"],
    ]
    for x in summary["source_type_stats"]
]
add_table(
    doc,
    ["来源类型", "记录", "URL", "有效", "部分", "不支持", "未核验"],
    source_rows,
    widths=[1.45, 0.65, 0.55, 0.55, 0.55, 0.65, 0.65],
    font_size=8.2,
)
add_heading(doc, "5.1 真实有效来源", 2)
add_bullet(doc, "品牌官网的江门区域页、门店页、案例页最能确认“品牌是否在江门提供相关服务”，但不能单独证明口碑、施工质量或售后表现。")
add_bullet(doc, "企业信息/百科可辅助确认主体身份、地址和经营范围；地图/点评页可辅助确认门店，但动态页面和反爬页常使正文无法核验。")
add_bullet(doc, "行业平台、媒体和内容平台只有在页面正文真正出现推荐品牌且与江门业务相符时才计入支持。")
add_heading(doc, "5.2 低质量与错配来源", 2)
add_bullet(doc, "软文/转载平台共118条推荐记录，其中87条保持未核验；页面常以“2026推荐”“口碑榜”“避坑指南”为标题，但缺少可追溯的评选方法。")
add_bullet(doc, "政府采购和江门政府页面虽权威，但多为公共工程或无关项目，不能据此支持住宅装修品牌推荐。")
add_bullet(doc, "同一品牌官网或软文页面被分配给多个品牌，是本次最常见的引用错配模式；高频引用不等于高证据质量。")

page_counts = Counter((x["citation_title"], x["citation_url"]) for x in rows if x["citation_url"])
top_pages = page_counts.most_common(12)
add_heading(doc, "5.3 最常被引用的页面", 2)
page_table = doc.add_table(rows=1, cols=3)
page_table.style = "Table Grid"
page_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["引用次数", "页面", "网站"]):
    set_cell_shading(page_table.rows[0].cells[i], BLUE)
    p = page_table.rows[0].cells[i].paragraphs[0]
    r = p.add_run(h)
    set_run_font(r, size=8.2, bold=True, color=WHITE)
set_repeat_table_header(page_table.rows[0])
for idx, ((title_text, url), count) in enumerate(top_pages):
    row = page_table.add_row()
    prevent_row_split(row)
    if idx % 2 == 0:
        for cell in row.cells:
            set_cell_shading(cell, "F7FAFC")
    p = row.cells[0].paragraphs[0]
    r = p.add_run(str(count))
    set_run_font(r, size=8.2, color=INK)
    p = row.cells[1].paragraphs[0]
    add_hyperlink(p, title_text[:55] + ("…" if len(title_text) > 55 else ""), url)
    p = row.cells[2].paragraphs[0]
    r = p.add_run(urlparse(url).netloc.lower())
    set_run_font(r, size=7.8, color=INK)
    for cell in row.cells:
        set_cell_margins(cell)
doc.add_paragraph().paragraph_format.space_after = Pt(0)

add_heading(doc, "官网依赖与第三方来源", 1)
depend_rows = []
for x in summary["overall"][:25]:
    if x["linked_count"] == 0:
        label = "无可核验链接"
    elif x["official_share"] >= 0.6:
        label = "主要依赖官网"
    elif x["official_share"] <= 0.2:
        label = "主要依赖第三方"
    else:
        label = "官网与第三方混合"
    depend_rows.append([
        x["brand"],
        x["linked_count"],
        x["official_count"],
        x["third_party_count"],
        f'{x["official_share"] * 100:.1f}%' if x["official_share"] is not None else "—",
        label,
        x["most_cited_sites"][0]["site"] if x["most_cited_sites"] else "",
    ])
add_table(
    doc,
    ["品牌", "带链接", "官网", "第三方", "官网占比", "依赖类型", "最常引用网站"],
    depend_rows,
    widths=[1.25, 0.62, 0.55, 0.62, 0.72, 1.15, 1.35],
    font_size=7.7,
)
add_callout(
    doc,
    "依赖结构观察",
    "索菲亚、华美乐、居众主要依赖官网；健威、峰尚汇、森之原、名雕主要依赖第三方。劳卡、华浔、星艺、名匠属于官网与第三方混合型，但部分第三方链接仍存在错配或软文问题。",
)

add_heading(doc, "楼盘场景", 1)
scene = summary["track_top10"]["楼盘场景"]
scene_rows = [
    [i + 1, x["brand"], x["count"], f'{x["average_rank"]:.2f}', x["brand_scope"]]
    for i, x in enumerate(scene)
]
add_table(
    doc,
    ["位次", "品牌", "推荐次数", "平均位次", "范围"],
    scene_rows,
    widths=[0.55, 1.55, 0.9, 0.9, 1.55],
    font_size=8.5,
)
add_body(
    doc,
    "楼盘场景中，欧派（11次）、索菲亚（10次）、华浔品味装饰（10次）居前。全国品牌更容易被AI用于“精装升级、柜体、家具、窗帘”的标准化表达；装修集团在毛坯大户型和全案交付问题中更常出现。",
)
add_body(
    doc,
    "但“被推荐”不等于“在该楼盘有经验”。六个楼盘问题中，公开引用很少提供可核验的同楼盘实景完工案例、具体栋号/户型、完工日期和业主评价；无法核验时均未推定品牌具有项目经验。",
)

add_heading(doc, "江门本地品牌 vs 全国连锁品牌", 1)
add_image(doc, chart_scope, width=5.9, caption="图4　品牌范围与AI推荐量")
scope_rows = [
    [
        x["scope"],
        x["recommendation_rows"],
        pct(x["recommendation_rows"], 757),
        x["unique_brands"],
        f'{x["average_rank"]:.2f}',
        x["unique_questions"],
        x["valid_or_partial"],
    ]
    for x in summary["scope_stats"]
]
add_table(
    doc,
    ["品牌范围", "推荐记录", "占比", "品牌数", "平均位次", "覆盖题次", "有效/部分"],
    scope_rows,
    widths=[1.35, 0.78, 0.65, 0.65, 0.75, 0.75, 0.85],
    font_size=8.3,
)
add_bullet(doc, "全国/跨区域品牌占518条推荐（68.4%），覆盖更广，主要由区域官网、总部官网和大量第三方页面支撑。")
add_bullet(doc, "江门本地品牌占211条推荐（27.9%），平均位次同为2.70；在高端定制、整木与红木细分场景更容易进入前两位。")
add_bullet(doc, "本地品牌的主要短板不是“没有被推荐”，而是官网结构化信息薄弱、第三方证据来源集中、楼盘与完工案例难以交叉核验。")

add_heading(doc, "当前GEO竞争最强的品牌", 1)
competitors = [
    ["华浔品味装饰", "72次 / 2.50", "装修、高端装修、全屋设计三赛道持续出现；江门区域官网可核验公司地址、整装/全案、设计师和案例。", "第三方内容占比高，仍有泛榜单与错配链接。"],
    ["星艺装饰", "64次 / 2.28", "江门站包含流程、案例、设计师、工地预约和售后表达，结构最完整之一。", "个别AI引用被其他品牌页面替代。"],
    ["名匠装饰", "56次 / 2.54", "江门/新会区域站和高端大宅关键词覆盖稳定。", "多域名、移动站和镜像页造成来源分散。"],
    ["健威家居", "51次 / 1.69", "本地属性、高端定制、家具与全屋设计语义强，平均位次最高。", "官网引用占比仅8.0%，主要依赖公众号、头条和百科。"],
    ["劳卡", "50次 / 2.92", "江门整装官网覆盖装修、定制、翻新与一站式交付，跨赛道能力强。", "同一官网被AI错配给其他品牌，且部分引用来自软文。"],
    ["索菲亚", "48次 / 2.38", "江门门店/攻略页搜索可见性强，官网占比79.4%，全屋定制与楼盘场景领先。", "官网内容自述较多，第三方江门交付案例不足。"],
    ["欧派", "47次 / 2.77", "全屋定制与精装升级语义稳定，楼盘场景推荐次数第一。", "本次有效证据偏弱，多条引用实际指向其他品牌或无关政府页面。"],
    ["峰尚汇装饰", "28次 / 2.64", "本地装修关键词曝光集中、单篇软文带来高频命中。", "20次引用集中于同一中华网页面，来源结构单一。"],
]
add_table(
    doc,
    ["品牌", "次数/均位", "内容优势", "信息风险"],
    competitors,
    widths=[1.1, 0.8, 2.4, 2.1],
    font_size=7.7,
)
add_heading(doc, "可见性优势的共同机制", 2)
add_bullet(doc, "有独立江门区域页，标题直接包含“江门+品牌+业务”。")
add_bullet(doc, "同时覆盖设计、案例、工艺、材料、门店、流程、工地、售后等多种可检索实体。")
add_bullet(doc, "官网与第三方共同存在，而不是只靠一篇榜单或单一公众号稿。")
add_bullet(doc, "内容能对应不同消费阶段：收楼、毛坯、精装升级、大平层、别墅、全案、软装。")

add_heading(doc, "公开内容的明显空白选题", 1)
gaps = [
    ["楼盘实景案例库", "海悦天玺、海悦天晟、蒲葵天心、博学名苑、潮闻东方、骏景湾天汇；按户型/面积/交付状态整理。", "每案含现场照片、户型图、完工日期、设计师、施工/安装节点和可参观状态。"],
    ["一体化服务边界", "设计、施工、柜体、橱柜、家具、窗帘、软装分别由谁交付。", "用责任矩阵和合同边界解释“能做”与“真正统筹”的差异。"],
    ["本地工厂与直营证明", "工厂地址、产线、设备、门店性质、服务半径。", "官网独立页链接工商主体、现场图、生产流程和预约参观方式。"],
    ["高端交付量化", "大平层/别墅的深化设计、节点收口、木作、石材、灯光、机电协同。", "公开节点标准、验收清单、样板段、交付周期与变更机制。"],
    ["环保可核验资料", "板材、涂料、胶黏剂、成品家具的检测与适用范围。", "发布检测报告原件、批次、有效期和对应产品，而非只写“环保”。"],
    ["安装与售后SLA", "安装团队、问题响应、复尺、补件、质保和维修。", "明确时效、责任人、质保边界和真实工单案例。"],
    ["真实工地与完工回访", "可参观工地、阶段照片、隐蔽工程、交付后回访。", "采用时间线形式，并允许用户核对项目状态。"],
    ["第三方证据", "行业协会、媒体、地图门店、业主评价与公开投诉处理。", "减少同稿多站转载，增加可追溯主体和方法说明。"],
    ["本地品牌对比全国连锁", "报价透明度、设计深度、工厂柔性、交付半径、售后稳定性。", "使用统一指标和真实样本，不做无方法的“Top10”。"],
    ["生活方式与空间规划", "家庭结构、收纳、动线、适老、亲子、宠物、家政和长期变化。", "用入住后复盘证明全屋设计不只是风格效果图。"],
]
add_table(
    doc,
    ["空白主题", "内容范围", "建议的可核验要素"],
    gaps,
    widths=[1.25, 2.25, 2.65],
    font_size=7.8,
)
add_callout(
    doc,
    "优先级判断",
    "最稀缺、也最能提升AI可信推荐的不是再做一篇“江门装修公司排名”，而是同楼盘真实案例、服务责任边界、工厂/门店证明、交付节点和售后SLA。这些内容可同时减少AI幻觉、错配引用和“未核验”比例。",
    fill="FFF8E7",
    accent=SAND,
)

add_heading(doc, "结论", 1)
add_body(
    doc,
    "当前江门AI搜索生态并非由单一市场领导者主导，而是由“全国装修集团的跨赛道覆盖”“全国定制品牌的官网SEO”“江门本地品牌的高端细分标签”共同构成。华浔、星艺、名匠在装修与设计端最强；索菲亚、欧派在定制与楼盘升级端最强；健威、盛世周木匠在本地高端定制语义上最有辨识度；劳卡通过江门整装本地站实现跨装修、定制和设计覆盖。",
)
add_body(
    doc,
    "但证据层仍有明显缺口：55%左右的推荐记录保持未核验，另有119条出现页面与品牌不匹配。对品牌而言，GEO竞争的下一阶段不是增加更多泛榜单，而是把江门门店、真实工地、完工案例、楼盘经验、工厂与交付标准做成结构清晰、可被AI直接引用且可由第三方复核的公开资料。",
)

add_heading(doc, "附录：数据文件与异常说明", 1)
add_heading(doc, "A. 完整明细", 2)
p = add_body(doc, "完整工作簿：", size=9.5)
add_external_file_link(p, "江门AI搜索可见性测试明细_180题.xlsx", OUT / "江门AI搜索可见性测试明细_180题.xlsx")
add_body(
    doc,
    "工作簿包含“总览、五赛道Top10、品牌统计、平台对比、楼盘场景、来源质量、180题明细、引用核验、方法与限制”9个工作表。每条推荐记录均保留平台、问题编号、问题、品牌、位次、理由、引用页面、URL、来源类型、发布时间、支持度、搜索日期和会话URL。",
    size=9.2,
)
add_heading(doc, "B. 平台输出异常", 2)
anomaly_rows = [
    ["DeepSeek", "36题", "推荐结果可保存，但样本未提供可打开的引用URL，因此141条推荐全部保持未核验。"],
    ["千问", "第29题", "未给出可核验的品牌推荐，记录为零推荐样本。"],
    ["千问", "第31题", "未给出可核验的品牌推荐，记录为零推荐样本。"],
    ["千问", "第18题", "推荐中出现麦格门窗、建雅摩托、恒锐石材、强立等与高端定制语义不匹配的实体，保留用于衡量相关性失败。"],
    ["DeepSeek", "第32题", "第一推荐为“索菲亚+如鱼得水”组合，拆分为两个品牌记录，因此该题产生6条品牌记录。"],
]
for platform, issue, anomaly in anomaly_rows:
    add_bullet(doc, f"{platform}｜{issue}：{anomaly}", bold_lead=f"{platform}｜{issue}：")

add_heading(doc, "C. 关键可核验网页示例", 2)
example_links = [
    ("华浔品味装饰江门区域页", "https://www.hxdec.com/website-branch/191.html"),
    ("星艺装饰江门公司官网", "http://jiangmen.xydec.com.cn/"),
    ("劳卡江门整装官网", "https://www.jmlaokazz.com/"),
    ("索菲亚江门定制攻略页", "https://www.suofeiya.com/gonglue/25367.html"),
    ("华美乐江门站", "https://www.homello.com/index.php?city=5"),
    ("盛世周木匠整装定制页面", "http://www.sszmj.com/content-24-101-1.html"),
    ("森润整木官网", "https://www.senrun-wood.com/"),
]
for label, url in example_links:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    add_hyperlink(p, label, url)
    r = p.add_run(f"  {url}")
    set_run_font(r, size=8.2, color=MUTED)

for sec in doc.sections:
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

doc.core_properties.title = "江门全屋定制、装修与全屋设计AI搜索基准报告"
doc.core_properties.subject = "五平台180次独立联网AI搜索可见性盲测"
doc.core_properties.author = "Codex GEO Research"
doc.core_properties.keywords = "江门, GEO, AI搜索可见性, 全屋定制, 装修, 全屋设计"

report_path = OUT / "江门全屋定制、装修与全屋设计AI搜索基准报告.docx"
doc.save(report_path)
print(json.dumps({"report": str(report_path), "sections": len(doc.sections), "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}, ensure_ascii=False, indent=2))
